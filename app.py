"""Email Bot - a small local web app for sending form letters through the Outlook app
on a Mac, or classic Outlook on Windows.

Run it with:  python app.py   (or double-click start.command on a Mac, start.bat on Windows,
or EmailBot.exe). It only listens on this computer (127.0.0.1); nobody else on the network
can reach it.

    python app.py --diagnose [--account ADDRESS] [--send-to ADDRESS]   check Windows Outlook step by step
    python app.py --check-bundle                                       check EmailBot.exe has everything
"""

import logging
import os
import socket
import sys
import threading
import time
import traceback
import webbrowser
from pathlib import Path

from flask import Flask, abort, jsonify, render_template, request

from emailbot import importer, merge, storage
from emailbot.outlook import Outlook, OutlookError

PORT = int(os.environ.get("EMAILBOT_PORT", "5050"))
DEBUG = bool(os.environ.get("EMAILBOT_DEBUG"))
log = logging.getLogger("emailbot")
MAX_ATTACHMENT_BYTES = 3 * 1024 * 1024  # keeps every email well under mail servers' size limits

# The page's files sit next to this file, or inside the unpacked bundle when running as EmailBot.exe.
RESOURCES = Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parent))
app = Flask(__name__, template_folder=str(RESOURCES / "templates"), static_folder=str(RESOURCES / "static"))
app.config["MAX_CONTENT_LENGTH"] = 25 * 1024 * 1024
outlook = Outlook()


@app.before_request
def local_only():
    # Block other websites from talking to this app through your browser
    # (DNS rebinding and cross-site form posts).
    host = (request.host or "").rsplit(":", 1)[0]
    if host not in ("127.0.0.1", "localhost"):
        abort(403)
    if request.path.startswith("/api/") and request.method != "GET":
        if request.headers.get("X-EmailBot") != "1":
            abort(403)


@app.errorhandler(OutlookError)
@app.errorhandler(ValueError)
def user_error(e):
    if isinstance(e, OutlookError):
        log.warning("%s", e)
    return jsonify({"error": str(e)}), 400


@app.get("/")
def index():
    return render_template("index.html")


# ---- workspace (letters + recipients) -----------------------------------

@app.get("/api/state")
def get_state():
    return jsonify(storage.load_workspace())


@app.put("/api/state")
def put_state():
    ws = request.get_json()
    if not isinstance(ws, dict) or "letters" not in ws or "recipients" not in ws:
        raise ValueError("Bad data.")
    # Attachment lists are managed by the server; keep what's on disk.
    current = {l["id"]: l for l in storage.load_workspace()["letters"]}
    for letter in ws["letters"]:
        letter["attachments"] = current.get(letter["id"], {}).get("attachments", [])
    storage.save_workspace(ws)
    return jsonify({"ok": True})


@app.post("/api/letters/new")
def letter_new():
    copy_id = (request.get_json(silent=True) or {}).get("copy_from")
    ws = storage.load_workspace()
    source = next((l for l in ws["letters"] if l["id"] == copy_id), None)
    letter = storage.new_letter(source)
    ws["letters"].append(letter)
    ws["current_letter_id"] = letter["id"]
    storage.save_workspace(ws)
    return jsonify(ws)


@app.post("/api/letters/<letter_id>/delete")
def letter_delete(letter_id):
    ws = storage.load_workspace()
    ws["letters"] = [l for l in ws["letters"] if l["id"] != letter_id]
    if not ws["letters"]:
        ws["letters"].append(storage.new_letter())
    if ws["current_letter_id"] == letter_id:
        ws["current_letter_id"] = ws["letters"][0]["id"]
    storage.save_workspace(ws)
    return jsonify(ws)


def _letter(ws, letter_id):
    letter = next((l for l in ws["letters"] if l["id"] == letter_id), None)
    if letter is None:
        raise ValueError("That letter no longer exists.")
    return letter


@app.post("/api/letters/<letter_id>/attachments")
def attachment_add(letter_id):
    ws = storage.load_workspace()
    letter = _letter(ws, letter_id)
    folder = storage.attachment_dir(letter_id)
    for f in request.files.getlist("files"):
        name = storage.safe_filename(f.filename)
        data = f.read()
        existing = sum((folder / n).stat().st_size for n in letter["attachments"]
                       if (folder / n).exists() and n != name)
        if existing + len(data) > MAX_ATTACHMENT_BYTES:
            raise ValueError("Attachments are limited to 3 MB in total per email.")
        (folder / name).write_bytes(data)
        if name not in letter["attachments"]:
            letter["attachments"].append(name)
    storage.save_workspace(ws)
    return jsonify({"attachments": letter["attachments"]})


@app.post("/api/letters/<letter_id>/attachments/remove")
def attachment_remove(letter_id):
    name = (request.get_json() or {}).get("name", "")
    ws = storage.load_workspace()
    letter = _letter(ws, letter_id)
    if name in letter["attachments"]:
        letter["attachments"].remove(name)
        (storage.attachment_dir(letter_id) / storage.safe_filename(name)).unlink(missing_ok=True)
    storage.save_workspace(ws)
    return jsonify({"attachments": letter["attachments"]})


@app.post("/api/import")
def import_recipients():
    if "file" in request.files:
        f = request.files["file"]
        return jsonify(importer.parse(filename=f.filename, data=f.read()))
    return jsonify(importer.parse(text=(request.get_json() or {}).get("text", "")))


# ---- preview & sending ---------------------------------------------------

def build_emails(letter_id, resend=False):
    ws = storage.load_workspace()
    letter = _letter(ws, letter_id)
    sent = storage.load_sent().get(letter_id, {})
    emails = []
    for row in ws["recipients"]["rows"]:
        if not row.get("include", True):
            continue
        email = merge.render(letter, row)
        email["already_sent"] = sent.get(email["to"].lower(), "")
        email["will_send"] = not email["has_errors"] and (resend or not email["already_sent"])
        emails.append(email)
    return letter, emails


@app.post("/api/preview")
def preview():
    body = request.get_json() or {}
    letter, emails = build_emails(body.get("letter_id"), body.get("resend", False))
    return jsonify({"letter": letter["name"], "attachments": letter["attachments"], "emails": emails})


class SendJob:
    def __init__(self):
        self.lock = threading.Lock()
        self.state = {"running": False, "total": 0, "done": 0, "results": [], "test": False}
        self.cancel = False

    def start(self, letter, emails, test):
        with self.lock:
            if self.state["running"]:
                raise ValueError("Emails are already being sent.")
            self.state = {"running": True, "total": len(emails), "done": 0, "results": [], "test": test}
            self.cancel = False
        threading.Thread(target=self._run, args=(letter, emails, test), daemon=True).start()

    def _run(self, letter, emails, test):
        folder = storage.attachment_dir(letter["id"])
        files = [folder / n for n in letter["attachments"] if (folder / n).exists()]
        delay = float(outlook.config()["delay_seconds"])
        for i, e in enumerate(emails):
            if self.cancel:
                break
            if i and delay:
                time.sleep(delay)
            try:
                outlook.send(e["to"], e["cc"], e["subject"], e["html"], e["text"], files)
                ok, err = True, ""
                if not test:
                    storage.mark_sent(letter["id"], e["to"])
            except OutlookError as ex:
                log.warning("Sending to %s failed: %s", e["to"], ex)
                ok, err = False, str(ex)
            except Exception as ex:  # noqa: BLE001
                log.exception("Sending to %s failed unexpectedly", e["to"])
                ok, err = False, str(ex)
            storage.log_send(letter["name"] + (" (test)" if test else ""), e["to"], e["subject"], ok, err)
            with self.lock:
                self.state["done"] += 1
                self.state["results"].append({"to": e["to"], "ok": ok, "error": err})
        with self.lock:
            self.state["running"] = False
            self.state["cancelled"] = self.cancel


job = SendJob()


@app.post("/api/send")
def send():
    body = request.get_json() or {}
    letter, emails = build_emails(body.get("letter_id"), body.get("resend", False))
    if body.get("test"):
        status = outlook.status()
        me = status["test_address"]
        if not status["ready"]:
            raise OutlookError(status["unavailable_reason"])
        if not me:
            raise OutlookError("Enter your own email address in step 1 so the test has somewhere to go.")
        chosen = next((e for e in emails if e["row_id"] == body.get("row_id")), None) or (emails or [None])[0]
        if not chosen:
            raise ValueError("Add at least one recipient so there's something to test.")
        test_email = dict(chosen, to=me, cc=[], subject="[TEST] " + chosen["subject"])
        job.start(letter, [test_email], test=True)
    else:
        if any(e["has_errors"] for e in emails):
            raise ValueError("Some emails still have problems. Fix them (shown in red) before sending.")
        to_send = [e for e in emails if e["will_send"]]
        if not to_send:
            raise ValueError("There's nobody left to send to.")
        status = outlook.status()
        if not status["ready"]:
            raise OutlookError(status["unavailable_reason"])
        job.start(letter, to_send, test=False)
    return jsonify(job.state)


@app.get("/api/send/status")
def send_status():
    with job.lock:
        return jsonify(job.state)


@app.post("/api/send/cancel")
def send_cancel():
    job.cancel = True
    return jsonify({"ok": True})


@app.get("/api/history")
def history():
    return jsonify(storage.read_log())


# ---- Outlook connection --------------------------------------------------

@app.get("/api/outlook")
def outlook_status():
    return jsonify(outlook.status())


@app.post("/api/outlook/settings")
def outlook_settings():
    body = request.get_json() or {}
    changes = {}
    if "my_email" in body:
        changes["my_email"] = body["my_email"].strip()
    if "account" in body:
        changes["account"] = body["account"].strip()
    if "delay_seconds" in body:
        changes["delay_seconds"] = max(0, min(60, float(body["delay_seconds"])))
    outlook.save_config(**changes)
    return jsonify(outlook.status())


@app.post("/api/outlook/check")
def outlook_check():
    result = outlook.check()
    for step in result["steps"]:
        if not step["ok"]:
            log.warning("Outlook check failed: %s: %s", step["name"], step["detail"])
    return jsonify(result)


def check_bundle():
    """Used by the Windows build: make sure EmailBot.exe contains everything it needs."""
    client = app.test_client()
    host = {"Host": "127.0.0.1"}
    checks = {
        "page loads": client.get("/", headers=host).status_code == 200,
        "script loads": client.get("/static/app.js", headers=host).status_code == 200,
        "styles load": client.get("/static/style.css", headers=host).status_code == 200,
    }
    try:
        import openpyxl  # noqa: F401
        checks["Excel support"] = True
    except ImportError:
        checks["Excel support"] = False
    if sys.platform == "win32":
        try:
            import win32com.client  # noqa: F401
            checks["Outlook support (pywin32)"] = True
        except ImportError:
            checks["Outlook support (pywin32)"] = False
    for name, ok in checks.items():
        print(f"  [{'ok' if ok else 'FAILED'}] {name}")
    return 0 if all(checks.values()) else 1


def serve():
    url = f"http://127.0.0.1:{PORT}"
    open_browser = not os.environ.get("EMAILBOT_NO_BROWSER")
    try:
        with socket.socket() as s:
            s.bind(("127.0.0.1", PORT))
    except OSError:
        # Most likely Email Bot is already running in another window.
        print(f"\n  Email Bot already seems to be running. Opening {url}\n"
              "  If the page doesn't load, close the other Email Bot window and try again.\n")
        if open_browser:
            webbrowser.open(url)
        return
    log_file = storage.DATA_DIR / "emailbot.log"
    storage.DATA_DIR.mkdir(parents=True, exist_ok=True)
    logging.basicConfig(
        level=logging.DEBUG if DEBUG else logging.INFO,
        format="%(asctime)s %(levelname)s %(message)s",
        handlers=[logging.StreamHandler(), logging.FileHandler(log_file, encoding="utf-8")],
    )
    print(f"\n  Email Bot is running at {url}\n  Keep this window open while you use it. Close it to stop.\n"
          f"  Problems are written here and to {log_file}\n")
    if not DEBUG:
        # Hide the per-request log and the "development server" warning, which only
        # alarm non-technical users. Errors still show.
        logging.getLogger("werkzeug").setLevel(logging.ERROR)
    if open_browser:
        threading.Timer(1.0, lambda: webbrowser.open(url)).start()
    app.run(host="127.0.0.1", port=PORT, threaded=True)


def main():
    # Windows consoles can't show every character (e.g. curly quotes); never crash over it.
    for stream in (sys.stdout, sys.stderr):
        if stream is not None and hasattr(stream, "reconfigure"):
            stream.reconfigure(errors="replace")
    args = sys.argv[1:]
    if args[:1] == ["--diagnose"]:
        from emailbot import diagnose_windows
        sys.argv = [sys.argv[0] + " --diagnose"] + args[1:]
        return diagnose_windows.main()
    if args[:1] == ["--check-bundle"]:
        return check_bundle()
    serve()
    return 0


if __name__ == "__main__":
    try:
        code = main()
    except Exception:  # noqa: BLE001
        traceback.print_exc()
        code = 1
    if getattr(sys, "frozen", False) and code:
        # A double-clicked EmailBot.exe closes its window on exit; keep the message readable.
        try:
            input("\nSomething went wrong (details above). Press Enter to close this window.")
        except EOFError:  # no keyboard attached, e.g. run by the GitHub build
            pass
    sys.exit(code)

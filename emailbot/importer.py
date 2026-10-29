"""Reads recipient lists from CSV / Excel files or text pasted from a spreadsheet."""

import csv
import io
from datetime import date, datetime

from .storage import new_id


def _cell(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime):
        return value.strftime("%B %-d, %Y") if value.time() == datetime.min.time() else str(value)
    if isinstance(value, date):
        return value.strftime("%B %-d, %Y")
    return str(value).strip()


def _table_from_text(text):
    text = text.lstrip("﻿")
    sample = text[:5000]
    if "\t" in sample:
        delimiter = "\t"
    else:
        try:
            delimiter = csv.Sniffer().sniff(sample, delimiters=",;").delimiter
        except csv.Error:
            delimiter = ","
    return list(csv.reader(io.StringIO(text), delimiter=delimiter))


def _table_from_xlsx(data):
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise ValueError(
            "Opening Excel files isn't available because part of Email Bot's setup is missing. "
            "Use “Paste from Excel” instead, or save the file as CSV and open that."
        ) from None

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.worksheets[0]
    return [list(r) for r in ws.iter_rows(values_only=True)]


def _kind(header):
    h = header.lower().replace("-", "").replace("_", " ").strip()
    if h in ("cc", "copy to", "cc email"):
        return "cc"
    if "email" in h or h in ("to", "address"):
        return "email"
    return None


def parse(filename=None, data=None, text=None):
    """Return {"columns": [...], "rows": [...]} ready for the recipient table."""
    if text is not None:
        table = _table_from_text(text)
    elif filename and filename.lower().endswith((".xlsx", ".xlsm")):
        table = _table_from_xlsx(data)
    else:
        raw = data or b""
        try:
            decoded = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            decoded = raw.decode("cp1252", errors="replace")
        table = _table_from_text(decoded)

    table = [[_cell(c) for c in row] for row in table]
    table = [row for row in table if any(row)]
    if not table:
        raise ValueError("The file or pasted text didn't contain any rows.")

    headers = table[0]
    width = max(len(r) for r in table)
    headers += [""] * (width - len(headers))
    headers = [h or f"Column {i + 1}" for i, h in enumerate(headers)]

    email_col = next((i for i, h in enumerate(headers) if _kind(h) == "email"), None)
    cc_col = next((i for i, h in enumerate(headers) if _kind(h) == "cc"), None)
    if email_col is None:
        # No header row with "email" in it: look for a column of addresses instead.
        for i in range(width):
            if sum("@" in (r[i] if i < len(r) else "") for r in table[1:]) > 0:
                email_col = i
                break
    if email_col is None:
        raise ValueError("Couldn't find an email column. Make sure one column is titled “Email”.")

    columns = [h for i, h in enumerate(headers) if i not in (email_col, cc_col)]
    rows = []
    for r in table[1:]:
        r = r + [""] * (width - len(r))
        rows.append({
            "id": new_id(),
            "include": True,
            "email": r[email_col],
            "cc": r[cc_col] if cc_col is not None else "",
            "fields": {h: r[i] for i, h in enumerate(headers) if i not in (email_col, cc_col)},
            "addendums": [],
            "note": "",
        })
    return {"columns": columns, "rows": rows}
